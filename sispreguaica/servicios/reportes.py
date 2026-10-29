"""Reportes de ejecución presupuestaria y su exportación a CSV y Excel."""

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import CERO, Cuenta, Ejercicio, PartidaGasto, PartidaIngreso
from .catalogos import agrupar
from .presupuesto import SALDOS_GASTO, SALDOS_INGRESO, recalcular

# Niveles de agrupación: segmentos del código de cuenta que se conservan.
NIVELES = {"partida": 2, "generica": 3, "especifica": 4, "detalle": 5}


@dataclass
class Fila:
    codigo: str
    denominacion: str
    valores: dict[str, Decimal] = field(default_factory=dict)


@dataclass
class Reporte:
    titulo: str
    columnas: list[tuple[str, str]]  # (clave, título)
    filas: list[Fila]
    totales: dict[str, Decimal]


COLUMNAS_GASTO = [
    ("asignado", "Asignado"),
    ("aumento", "Aumentos"),
    ("disminucion", "Disminuciones"),
    ("actualizado", "Actualizado"),
    ("precomprometido", "Precomprometido"),
    ("comprometido", "Comprometido"),
    ("causado", "Causado"),
    ("pagado", "Pagado"),
    ("disponible", "Disponible"),
    ("por_pagar", "Por pagar"),
]
COLUMNAS_INGRESO = [
    ("previsto", "Previsto"),
    ("aumento", "Aumentos"),
    ("disminucion", "Disminuciones"),
    ("actualizado", "Actualizado"),
    ("devengado", "Devengado"),
    ("cobrado", "Cobrado"),
    ("por_devengar", "Por devengar"),
    ("por_cobrar", "Por cobrar"),
]


def _derivar_gasto(v: dict[str, Decimal]) -> None:
    v["actualizado"] = v["asignado"] + v["aumento"] - v["disminucion"]
    v["disponible"] = v["actualizado"] - v["precomprometido"] - v["comprometido"]
    v["por_pagar"] = v["causado"] - v["pagado"]


def _derivar_ingreso(v: dict[str, Decimal]) -> None:
    v["actualizado"] = v["previsto"] + v["aumento"] - v["disminucion"]
    v["por_devengar"] = v["actualizado"] - v["devengado"]
    v["por_cobrar"] = v["devengado"] - v["cobrado"]


def _denominaciones(sesion: Session) -> dict[str, str]:
    return dict(sesion.execute(select(Cuenta.codigo, Cuenta.denominacion)).all())


def ejecucion_gastos(
    sesion: Session,
    ejercicio: Ejercicio,
    nivel: str = "detalle",
    hasta: date | None = None,
    estructura: str = "",
) -> Reporte:
    gastos, _ = recalcular(sesion, ejercicio, hasta)
    nombres = _denominaciones(sesion)
    segmentos = NIVELES.get(nivel, 5)
    grupos: dict[tuple, Fila] = {}
    consulta = select(PartidaGasto).where(PartidaGasto.ejercicio_id == ejercicio.id)
    for p in sesion.scalars(consulta):
        if estructura and not p.estructura.codigo.startswith(estructura):
            continue
        cuenta = agrupar(p.cuenta.codigo, segmentos)
        if nivel == "detalle":
            clave = (p.estructura.codigo, cuenta, p.fuente.codigo)
            codigo = f"{p.estructura.codigo} {cuenta} F{p.fuente.codigo}"
        else:
            clave = (cuenta,)
            codigo = cuenta
        fila = grupos.setdefault(
            clave, Fila(codigo, nombres.get(cuenta, ""), dict.fromkeys(SALDOS_GASTO, CERO))
        )
        for saldo, monto in gastos.get(p.id, {}).items():
            fila.valores[saldo] += monto
    filas = [grupos[k] for k in sorted(grupos)]
    totales = dict.fromkeys(SALDOS_GASTO, CERO)
    for f in filas:
        for saldo in SALDOS_GASTO:
            totales[saldo] += f.valores[saldo]
        _derivar_gasto(f.valores)
    _derivar_gasto(totales)
    corte = f" al {hasta:%d/%m/%Y}" if hasta else ""
    return Reporte(
        f"Ejecución del presupuesto de gastos {ejercicio.anio}{corte}",
        COLUMNAS_GASTO,
        filas,
        totales,
    )


def ejecucion_ingresos(
    sesion: Session, ejercicio: Ejercicio, nivel: str = "detalle", hasta: date | None = None
) -> Reporte:
    _, ingresos = recalcular(sesion, ejercicio, hasta)
    nombres = _denominaciones(sesion)
    segmentos = NIVELES.get(nivel, 5)
    grupos: dict[str, Fila] = {}
    consulta = select(PartidaIngreso).where(PartidaIngreso.ejercicio_id == ejercicio.id)
    for p in sesion.scalars(consulta):
        cuenta = agrupar(p.cuenta.codigo, segmentos)
        fila = grupos.setdefault(
            cuenta, Fila(cuenta, nombres.get(cuenta, ""), dict.fromkeys(SALDOS_INGRESO, CERO))
        )
        for saldo, monto in ingresos.get(p.id, {}).items():
            fila.valores[saldo] += monto
    filas = [grupos[k] for k in sorted(grupos)]
    totales = dict.fromkeys(SALDOS_INGRESO, CERO)
    for f in filas:
        for saldo in SALDOS_INGRESO:
            totales[saldo] += f.valores[saldo]
        _derivar_ingreso(f.valores)
    _derivar_ingreso(totales)
    corte = f" al {hasta:%d/%m/%Y}" if hasta else ""
    return Reporte(
        f"Ejecución del presupuesto de ingresos {ejercicio.anio}{corte}",
        COLUMNAS_INGRESO,
        filas,
        totales,
    )


def a_csv(reporte: Reporte) -> bytes:
    salida = io.StringIO()
    w = csv.writer(salida, delimiter=";")
    w.writerow(["Código", "Denominación"] + [t for _, t in reporte.columnas])
    for f in reporte.filas:
        w.writerow(
            [f.codigo, f.denominacion]
            + [str(f.valores[c]).replace(".", ",") for c, _ in reporte.columnas]
        )
    w.writerow(["TOTAL", ""] + [str(reporte.totales[c]).replace(".", ",") for c, _ in reporte.columnas])
    # Con BOM para que Excel en español reconozca los acentos.
    return ("﻿" + salida.getvalue()).encode("utf-8")


def a_excel(reporte: Reporte, ente: str) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ejecución"
    hoja["A1"] = ente
    hoja["A1"].font = Font(bold=True, size=12)
    hoja["A2"] = reporte.titulo
    hoja["A2"].font = Font(bold=True)
    encabezado = ["Código", "Denominación"] + [t for _, t in reporte.columnas]
    hoja.append([])
    hoja.append(encabezado)
    azul = PatternFill("solid", fgColor="003366")
    for celda in hoja[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = azul
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
    formato = "#,##0.00"
    for f in reporte.filas:
        hoja.append([f.codigo, f.denominacion] + [f.valores[c] for c, _ in reporte.columnas])
    hoja.append(["TOTAL", ""] + [reporte.totales[c] for c, _ in reporte.columnas])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)
    for fila in hoja.iter_rows(min_row=5, min_col=3):
        for celda in fila:
            celda.number_format = formato
    hoja.column_dimensions["A"].width = 32
    hoja.column_dimensions["B"].width = 48
    for i in range(3, 3 + len(reporte.columnas)):
        hoja.column_dimensions[hoja.cell(row=4, column=i).column_letter].width = 17
    hoja.freeze_panes = "C5"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
